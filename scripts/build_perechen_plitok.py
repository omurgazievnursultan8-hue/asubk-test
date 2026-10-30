"""Собирает перечень плиток дашборда к финальному акту — приложение к ТЗ 24.

Источник один: docs/perechen-plitok/plitki.json. Руками правится только он, xlsx
пересобирается этим скриптом.

    python3 scripts/build_perechen_plitok.py                         # пути по умолчанию
    python3 scripts/build_perechen_plitok.py <plitki.json> <папка вывода>

Что собирается (в папке вывода, по умолчанию docs/perechen-plitok/):
- «Перечень плиток.xlsx», листы:
  «Плитки» — правила, обозначения и строки ПЛ-NN: где стоит (роль · секция · место), охват,
  отчёт или очередь, значения параметров, главное и второе число, сдвиг, светофор, виды по
  умолчанию, что утверждает заказчик;
  «Отчёты плиток» — строки ОП-NN: объект, отбор, показатели, разрезы, параметры и плитки,
  которые на отчёте стоят. Номер отчёта в листе «Плитки» — ссылка на его строку.

Вычисляется, а не пишется в JSON: плитки у отчёта, итоги в шапке.

Проверки (не сошлось — ничего не пишется):
- номера ПЛ-NN и ОП-NN идут подряд с 01 и не повторяются;
- у плитки заполнены все поля; обзорная ссылается на отчёт перечня, ссылочная — на очередь
  модуля; у каждого отчёта есть хотя бы одна плитка;
- «⚠»: «состав и порядок раскладки» — у плитки раскладки, «плитка каталога» — у плитки только
  в каталоге, «порог светофора» — ровно у плиток с порогом (знак > или <);
- в тексте нет внутренних кодов (ADR-, ДБ-, ОД-…) и номеров задач;
- сверка с макетом mockups/dashboard/dashboard.html (scripts/perechen_plitok_maket.mjs): те же
  плитки каталога с теми же названиями, видом, отчётом и очередью; второе число есть там же;
  светофор — те же цвета, знаки и границы, зелёный и число, на котором горит; та же база
  сдвига; те же виды по умолчанию (по роду вида: Движение, Разрез, Ряд, Карта, Таблица); те же
  значения параметров; те же отчёты под плитками; раскладки ролей — те же плитки в тех же
  секциях и в том же порядке.
  Не сверяются (тексты перечня): модуль, охват роли, слова главного и второго числа, разрез и
  колонки внутри вида.
"""
import json
import re
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / 'docs/perechen-plitok/plitki.json'
OUT_DIR = ROOT / 'docs/perechen-plitok'
OUT_NAME = 'Перечень плиток.xlsx'
MAKET = ROOT / 'scripts/perechen_plitok_maket.mjs'

TILE_RE = re.compile(r'^ПЛ-\d{2}$')
REPORT_RE = re.compile(r'^ОП-\d{2}$')
KINDS = ('обзорная', 'ссылочная')
TILE_REQ = ('maket', 'name', 'module', 'kind', 'scope', 'params', 'main', 'second', 'shift', 'tl',
            'views', 'warn')
REPORT_REQ = ('maket', 'name', 'object', 'filter')
WARN_LAYOUT, WARN_CATALOG, WARN_TL = 'состав и порядок раскладки', 'плитка каталога', 'порог светофора'
FORBID = re.compile(r'\b(?:ADR|ИО|ПВ|ОЧ|ОД|КВ|СС|РЯ|ДБ|ИМ|ГП|КН|ВБ|ПЗ)-Д?\d|\b[PП]\d+-R\d')
# база сдвига макета → слова перечня; у ссылочной плитки сдвига нет («—»)
SHIFTS = {'prevMonthEnd': 'за месяц', 'quarterStart': 'за квартал', 'halfStart': 'за полугодие',
          'yearStart': 'за год', 'yearAgo': 'за 12 месяцев', 'period': 'к прошлому периоду',
          'none': 'без сдвига'}
VIEWS = {'move': 'Движение', 'slice': 'Разрез', 'series': 'Ряд', 'map': 'Карта', 'table': 'Таблица'}
TL_COLORS = {'red': 'красный', 'yellow': 'жёлтый'}
# светофор на втором числе: «на втором: …», «на доле: …», у «Мои задания» — «просрочено > 0»
TL_ON_SECOND = re.compile(r'^на (?:втором|доле)\b|просрочено')

COLS = [
    ('Номер', 8), ('Плитка', 32), ('Модуль', 15), ('Вид', 11), ('Где стоит', 34), ('Охват роли', 20),
    ('Отчёт или очередь', 36), ('Значения параметров', 26), ('Главное число', 24),
    ('Второе число', 20), ('Сдвиг', 14), ('Светофор', 22), ('Виды по умолчанию', 38),
    ('⚠ Утверждает заказчик', 22),
]
REPORT_COLS = [
    ('Номер', 8), ('Отчёт', 34), ('Объект', 22), ('Отбор', 34), ('Показатели', 40),
    ('Разрезы', 36), ('Параметры', 22), ('Плитки на отчёте', 30),
]
REPORTS_SHEET = 'Отчёты плиток'

BLUE = '006AF5'
HEAD_FILL = PatternFill('solid', fgColor=BLUE)
WARN_FILL = PatternFill('solid', fgColor='FFF4D6')
THIN = Side(style='thin', color='C9D3E0')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical='top')


def fail(errors):
    for e in errors:
        print('✗', e, file=sys.stderr)
    print(f'ошибок: {len(errors)} — файл не записан', file=sys.stderr)
    sys.exit(1)


def strings(x):
    if isinstance(x, str):
        yield x
    elif isinstance(x, dict):
        for v in x.values():
            yield from strings(v)
    elif isinstance(x, list):
        for v in x:
            yield from strings(v)


def check_numbers(items, rx, prefix, errors):
    for i, it in enumerate(items, 1):
        want = f'{prefix}-{i:02d}'
        if not rx.match(str(it.get('id'))) or it['id'] != want:
            errors.append(f'{it.get("id")}: ожидался номер {want} — номера идут подряд с 01')


def check_tiles(data, errors):
    reports = {r['id'] for r in data['reports']}
    for t in data['tiles']:
        tid = t.get('id')
        for k in TILE_REQ:
            if not isinstance(t.get(k), str) or not t[k].strip():
                errors.append(f'{tid}: нет поля «{k}»')
        if t.get('kind') not in KINDS:
            errors.append(f'{tid}: вид «{t.get("kind")}» — обзорная или ссылочная')
        src = t.get('source') or {}
        if t.get('kind') == 'обзорная' and src.get('report') not in reports:
            errors.append(f'{tid}: обзорная плитка ссылается на отчёт перечня, а не на {src}')
        if t.get('kind') == 'ссылочная' and not (src.get('queue') and src.get('module')):
            errors.append(f'{tid}: ссылочная плитка ссылается на очередь модуля (queue, module)')
        where = t.get('where')
        if not isinstance(where, list) or any(
                not (w.get('role') and w.get('section') and isinstance(w.get('n'), int)) for w in where):
            errors.append(f'{tid}: «где стоит» — список {{role, section, n}}')
            where = []
        warn = [w.strip() for w in t.get('warn', '').split(',')]
        if bool(where) != (WARN_LAYOUT in warn) or (not where) != (WARN_CATALOG in warn):
            errors.append(f'{tid}: ⚠ «{WARN_LAYOUT}» у плитки раскладки, «{WARN_CATALOG}» — у плитки '
                          f'только в каталоге')
        if bool(re.search(r'[<>]', t.get('tl', ''))) != (WARN_TL in warn):
            errors.append(f'{tid}: ⚠ «{WARN_TL}» ставится ровно у плиток с порогом')
        bad = set(warn) - {WARN_LAYOUT, WARN_CATALOG, WARN_TL}
        if bad:
            errors.append(f'{tid}: в ⚠ неизвестное {sorted(bad)}')


def check_reports(data, errors):
    used = {t['source'].get('report') for t in data['tiles'] if t.get('source')}
    for r in data['reports']:
        for k in REPORT_REQ:
            if not isinstance(r.get(k), str) or not r[k].strip():
                errors.append(f'{r.get("id")}: нет поля «{k}»')
        for k in ('cols', 'params'):
            if not isinstance(r.get(k), list) or not r[k]:
                errors.append(f'{r.get("id")}: «{k}» — непустой список')
        if not isinstance(r.get('dims'), list):
            errors.append(f'{r.get("id")}: «dims» — список')
        if r.get('id') not in used:
            errors.append(f'{r.get("id")}: на отчёте не стоит ни одна плитка')


def tl_diff(text, tl):
    """Расхождение текста светофора перечня с порогом макета; None — сошлось."""
    if not tl:
        return None if text == '—' else 'в макете светофора нет'
    if text == '—':
        return 'в макете светофор есть'
    have = {c: (s, n) for c, s, n in re.findall(r'(жёлтый|красный)[^,;]*?([<>])\s*(\d+)', text)}
    want = {TL_COLORS[c]: (tl[c][0], str(tl[c][1])) for c in TL_COLORS if tl.get(c)}
    if have != want:
        return f'пороги {have}, в макете {want}'
    if ('зелёный' in text) != bool(tl.get('green')):
        return 'зелёный ' + ('есть' if tl.get('green') else 'не показывается') + ' в макете'
    if bool(TL_ON_SECOND.search(text)) != (tl.get('on') == 'second'):
        return f'светофор в макете на {"втором" if tl.get("on") == "second" else "главном"} числе'
    return None


def view_kinds(text):
    return [] if text == '—' else sorted(v.strip().split()[0] for v in text.split(';'))


def check_tile_maket(t, x, errors):
    tid = t['id']
    if (t['second'] != '—') != x['second']:
        errors.append(f'{tid}: второе число {"есть" if x["second"] else "нет"} в макете')
    why = tl_diff(t['tl'], x['tl'])
    if why:
        errors.append(f'{tid}: светофор «{t["tl"]}» — {why}')
    shift = SHIFTS.get(x['shift'], x['shift']) if x['shift'] else '—'
    if t['shift'] != shift:
        errors.append(f'{tid}: сдвиг «{t["shift"]}», в макете «{shift}»')
    want = sorted(VIEWS[k] for k in x['views'])
    if view_kinds(t['views']) != want:
        errors.append(f'{tid}: виды {view_kinds(t["views"])}, в макете {want}')
    params = x['params'] or '—'
    if t['params'] != params:
        errors.append(f'{tid}: параметры «{t["params"]}», в макете «{params}»')


def check_maket(data, errors):
    try:
        run = subprocess.run(['node', str(MAKET)], capture_output=True, text=True)
    except OSError as e:
        errors.append(f'сверка с макетом не запустилась: {e}')
        return
    if run.returncode:
        errors.append(f'сверка с макетом не запустилась: {run.stderr.strip() or run.returncode}')
        return
    mk = json.loads(run.stdout)
    mt = {t['id']: t for t in mk['tiles']}
    rep = {r['id']: r['maket'] for r in data['reports']}
    ours = {t['maket']: t for t in data['tiles']}
    for m in sorted(set(mt) - set(ours)):
        errors.append(f'плитки каталога макета «{mt[m]["title"]}» ({m}) нет в перечне')
    for m in sorted(set(ours) - set(mt)):
        errors.append(f'{ours[m]["id"]}: в каталоге макета нет плитки {m}')
    for m, t in ours.items():
        x = mt.get(m)
        if not x:
            continue
        if t['name'] != x['title']:
            errors.append(f'{t["id"]}: название «{t["name"]}», в макете «{x["title"]}»')
        if (t['kind'] == 'ссылочная') != (x['kind'] == 'link'):
            errors.append(f'{t["id"]}: вид «{t["kind"]}», в макете {x["kind"]}')
        src = t.get('source') or {}
        if x['report'] and rep.get(src.get('report')) != x['report']:
            errors.append(f'{t["id"]}: отчёт {src.get("report")}, в макете {x["report"]}')
        if x['queue'] and src.get('queue') != x['queue']:
            errors.append(f'{t["id"]}: очередь «{src.get("queue")}», в макете «{x["queue"]}»')
        check_tile_maket(t, x, errors)
    mr = {r['id']: r['title'] for r in mk['reports']}
    for r in data['reports']:
        if r['maket'] not in mr:
            errors.append(f'{r["id"]}: в макете под плитками каталога нет отчёта {r["maket"]}')
        elif r['name'] != mr[r['maket']]:
            errors.append(f'{r["id"]}: название «{r["name"]}», в макете «{mr[r["maket"]]}»')
    for m in sorted(set(mr) - set(rep.values())):
        errors.append(f'отчёта макета «{mr[m]}» ({m}) нет в перечне')
    placed = {}
    for t in data['tiles']:
        for w in t['where']:
            placed.setdefault(w['role'], []).append((w['n'], w['section'], t['maket']))
    for lay in mk['layouts']:
        want = [(i, x['section'], x['tile']) for i, x in enumerate(lay['items'], 1)]
        have = sorted(placed.pop(lay['role'], []))
        if have != want:
            diff = next((f'место {a[0]}: в перечне {a[1:]}, в макете {b[1:]}'
                         for a, b in zip(have, want) if a != b),
                        f'плиток в перечне {len(have)}, в макете {len(want)}')
            errors.append(f'раскладка «{lay["role"]}» расходится с макетом — {diff}')
    for role in placed:
        errors.append(f'роли «{role}» нет среди объявленных раскладок макета')


def check(data):
    errors = []
    for k in ('title', 'date', 'stage', 'rules', 'legend', 'tiles', 'reports'):
        if not data.get(k):
            errors.append(f'нет «{k}»')
    if errors:
        fail(errors)
    check_numbers(data['tiles'], TILE_RE, 'ПЛ', errors)
    check_numbers(data['reports'], REPORT_RE, 'ОП', errors)
    check_tiles(data, errors)
    check_reports(data, errors)
    for s in strings(data):
        m = FORBID.search(s)
        if m:
            errors.append(f'внутренний код «{m.group(0)}» в «{s[:60]}…»')
    if errors:
        fail(errors)
    check_maket(data, errors)
    if errors:
        fail(errors)


# --- xlsx --------------------------------------------------------------------------------

def set_widths(ws, cols):
    for i, (_, w) in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def text_row(ws, row, text, ncols, font=None, height=None):
    ws.cell(row=row, column=1, value=text)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    c = ws.cell(row=row, column=1)
    c.alignment = WRAP
    if font:
        c.font = font
    if height:
        ws.row_dimensions[row].height = height


def header(ws, row, cols):
    for i, (title, _) in enumerate(cols, 1):
        c = ws.cell(row=row, column=i, value=title)
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical='center')
        c.border = BORDER
    ws.row_dimensions[row].height = 30


def table_row(ws, r, vals):
    for i, v in enumerate(vals, 1):
        c = ws.cell(row=r, column=i, value=v)
        c.alignment = WRAP
        c.border = BORDER
    ws.cell(row=r, column=1).font = Font(bold=True)


def page_setup(ws, head):
    ws.print_title_rows = f'{head}:{head}'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def where_text(t):
    if not t['where']:
        return 'только каталог'
    return '\n'.join(f'{w["role"]} · {w["section"]} · {w["n"]}' for w in t['where'])


def source_text(t, names):
    s = t['source']
    if 'report' in s:
        return f'{s["report"]} «{names[s["report"]]}»'
    return f'очередь «{s["queue"]}», модуль «{s["module"]}»'


def sheet_tiles(wb, data, report_rows):
    ws = wb.active
    ws.title = 'Плитки'
    nc = len(COLS)
    set_widths(ws, COLS)
    tiles = data['tiles']
    names = {r['id']: r['name'] for r in data['reports']}
    link = sum(1 for t in tiles if t['kind'] == 'ссылочная')
    placed = sum(1 for t in tiles if t['where'])
    r = 1
    text_row(ws, r, data['title'], nc, Font(bold=True, size=14), 22)
    r += 1
    text_row(ws, r, f'{data["date"]}. {data["stage"]}', nc, Font(italic=True, color='555555'))
    r += 1
    total = (f'Плиток: {len(tiles)} — обзорных {len(tiles) - link}, ссылочных {link}; в раскладках '
             f'ролей {placed}, только в каталоге {len(tiles) - placed}. Отчётов плиток: '
             f'{len(data["reports"])} (лист «{REPORTS_SHEET}»).')
    text_row(ws, r, total, nc, Font(bold=True))
    r += 2
    for rule in data['rules']:
        text_row(ws, r, rule, nc, height=30)
        r += 1
    r += 1
    text_row(ws, r, 'Обозначения: ' + '; '.join(data['legend']) + '.', nc,
             Font(color='555555', size=9), 30)
    r += 2
    header(ws, r, COLS)
    head = r
    for t in tiles:
        r += 1
        table_row(ws, r, [t['id'], t['name'], t['module'], t['kind'], where_text(t), t['scope'],
                          source_text(t, names), t['params'], t['main'], t['second'], t['shift'],
                          t['tl'], t['views'], t['warn']])
        rep = t['source'].get('report')
        if rep:
            c = ws.cell(row=r, column=7)
            c.hyperlink = f"#'{REPORTS_SHEET}'!A{report_rows[rep]}"
            c.font = Font(color=BLUE, underline='single')
        ws.cell(row=r, column=nc).fill = WARN_FILL
    ws.freeze_panes = ws.cell(row=head + 1, column=3)
    ws.auto_filter.ref = f'A{head}:{get_column_letter(nc)}{r}'
    page_setup(ws, head)


def report_layout(data):
    """Строка каждого отчёта на листе «Отчёты плиток» — нужна ссылкам с листа «Плитки»."""
    head = 4
    return head, {rep['id']: head + i for i, rep in enumerate(data['reports'], 1)}


def sheet_reports(wb, data):
    ws = wb.create_sheet(REPORTS_SHEET)
    nc = len(REPORT_COLS)
    set_widths(ws, REPORT_COLS)
    on = {}
    for t in data['tiles']:
        rep = t['source'].get('report')
        if rep:
            on.setdefault(rep, []).append(t['id'])
    text_row(ws, 1, f'{data["title"]}: отчёты плиток', nc, Font(bold=True, size=14), 22)
    text_row(ws, 2, 'Отчёт плитки — общий отчёт каталога: объект, отбор, показатели, разрезы и '
                    'параметры. Плитка берёт из него главное и второе число и задаёт значения '
                    'параметров. В перечень отчётов кредитного модуля отчёты плиток не входят.',
             nc, Font(italic=True, color='555555'), 30)
    head, rows = report_layout(data)
    header(ws, head, REPORT_COLS)
    r = head
    for rep in data['reports']:
        r += 1
        assert r == rows[rep['id']]
        table_row(ws, r, [rep['id'], rep['name'], rep['object'], rep['filter'],
                          '; '.join(rep['cols']), '; '.join(rep['dims']) or '—',
                          '; '.join(rep['params']), ', '.join(on[rep['id']])])
    ws.freeze_panes = ws.cell(row=head + 1, column=3)
    ws.auto_filter.ref = f'A{head}:{get_column_letter(nc)}{r}'
    page_setup(ws, head)


def main():
    args = sys.argv[1:]
    src = Path(args[0]) if args else SRC
    out = Path(args[1]) if len(args) > 1 else OUT_DIR
    data = json.loads(src.read_text(encoding='utf-8'))
    check(data)
    wb = Workbook()
    sheet_tiles(wb, data, report_layout(data)[1])
    sheet_reports(wb, data)
    out.mkdir(parents=True, exist_ok=True)
    path = out / OUT_NAME
    wb.save(path)
    print(f'{path.relative_to(ROOT) if path.is_relative_to(ROOT) else path}: '
          f'плиток {len(data["tiles"])}, отчётов {len(data["reports"])}; сверка с макетом — сошлось')


if __name__ == '__main__':
    main()
